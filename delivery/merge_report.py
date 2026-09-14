#!/usr/bin/env python3
"""server_dbms_result.xlsx + cloud_result.xlsx -> infra_check_{YYYYMMDD}_{회차}.xlsx

실행 예시:
    python3 merge_report.py --round "1차" --date 20260914

시트 4개(요약/서버/DBMS/클라우드)로 구성한 최종 결과물을 만든다. 이 파일이 SHA-256
해시와 함께 보안팀에 전달되는 결과물이다.
"""
import argparse
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

COLUMNS = ["항목ID", "항목명", "판정", "상세", "대상", "점검일시", "회차"]

STATUS_FILL = {
    "PASS": PatternFill("solid", fgColor="C6EFCE"),
    "FAIL": PatternFill("solid", fgColor="FFC7CE"),
    "SKIP": PatternFill("solid", fgColor="D9D9D9"),
    "N/A": PatternFill("solid", fgColor="F2F2F2"),
    "REVIEW": PatternFill("solid", fgColor="E4DFEC"),
}
STATUSES = ["PASS", "FAIL", "N/A", "SKIP", "REVIEW"]


def parse_args():
    script_dir = Path(__file__).resolve().parent
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--date", default=datetime.now().strftime("%Y%m%d"),
                   help="날짜 폴더명(YYYYMMDD, 기본: 오늘) — script/results/{date}/ 아래에서 입출력")
    p.add_argument("--server-dbms", default=None,
                   help="입력 경로(기본: script/results/{date}/server_dbms_result.xlsx)")
    p.add_argument("--cloud", default=None,
                   help="입력 경로(기본: script/results/{date}/cloud_result.xlsx)")
    p.add_argument("--round", dest="round_", default="정기점검")
    p.add_argument("--output", default=None,
                   help="출력 경로(기본: script/results/{date}/infra_check_{date}_{round}.xlsx)")
    args = p.parse_args()
    results_dir = script_dir / "results" / args.date
    if args.server_dbms is None:
        args.server_dbms = str(results_dir / "server_dbms_result.xlsx")
    if args.cloud is None:
        args.cloud = str(results_dir / "cloud_result.xlsx")
    if args.output is None:
        args.output = str(results_dir / f"infra_check_{args.date}_{args.round_}.xlsx")
    return args


def read_sheet_rows(path: Path, sheet_name: str):
    if not path.exists():
        print(f"[경고] {path} 없음 — {sheet_name} 시트 비워둠")
        return []
    wb = load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        print(f"[경고] {path}에 '{sheet_name}' 시트 없음")
        return []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return [list(r) for r in rows if r and r[0]]


def write_data_sheet(wb, title, rows):
    ws = wb.create_sheet(title=title)
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
        status = row[2]
        fill = STATUS_FILL.get(status)
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill
    for col_idx, width in enumerate([10, 46, 10, 70, 20, 26, 12], start=1):
        ws.column_dimensions[chr(64 + col_idx)].width = width
    ws.freeze_panes = "A2"
    return ws


def write_summary_sheet(wb, round_, checked_at, all_rows):
    ws = wb.create_sheet(title="요약")
    ws.append(["구분", "값"])
    ws["A1"].font = ws["B1"].font = Font(bold=True)
    counts = {s: sum(1 for r in all_rows if r[2] == s) for s in STATUSES}
    summary_data = [
        ("점검일시", checked_at),
        ("회차", round_),
        ("전체 항목 수", len(all_rows)),
        ("PASS", counts["PASS"]),
        ("FAIL", counts["FAIL"]),
        ("N/A", counts["N/A"]),
        ("SKIP", counts["SKIP"]),
        ("REVIEW", counts["REVIEW"]),
    ]
    for label, value in summary_data:
        ws.append([label, value])
        if label in STATUS_FILL:
            ws.cell(row=ws.max_row, column=1).fill = STATUS_FILL[label]
            ws.cell(row=ws.max_row, column=2).fill = STATUS_FILL[label]
    if counts["SKIP"] > 0:
        ws.append([])
        note = ws.cell(row=ws.max_row + 1, column=1,
                        value=f"⚠ SKIP {counts['SKIP']}건 — 기준값 미확정으로 판정 불가한 항목입니다. "
                              "group_vars/all.yml, cloud/config.py 값 확정 후 재점검 필요.")
        note.font = Font(color="C00000", bold=True)
        ws.merge_cells(start_row=note.row, start_column=1, end_row=note.row, end_column=2)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 60
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    return ws


def main():
    args = parse_args()
    server_rows = read_sheet_rows(Path(args.server_dbms), "서버")
    dbms_rows = read_sheet_rows(Path(args.server_dbms), "DBMS")
    cloud_rows = read_sheet_rows(Path(args.cloud), "클라우드")
    all_rows = server_rows + dbms_rows + cloud_rows

    if not all_rows:
        raise SystemExit("병합할 데이터가 없습니다 — server_dbms_result.xlsx / cloud_result.xlsx 경로를 확인하세요.")

    checked_at = max((r[5] for r in all_rows if r[5]), default=datetime.now().isoformat())

    wb = Workbook()
    wb.remove(wb.active)
    write_summary_sheet(wb, args.round_, checked_at, all_rows)
    write_data_sheet(wb, "서버", server_rows)
    write_data_sheet(wb, "DBMS", dbms_rows)
    write_data_sheet(wb, "클라우드", cloud_rows)

    out_path = Path(args.output)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"[완료] {out_path} 생성 — 총 {len(all_rows)}건 "
          f"(서버 {len(server_rows)} / DBMS {len(dbms_rows)} / 클라우드 {len(cloud_rows)})")
    print("[다음 단계] mgmt 서버에서 SHA-256 해시 계산 필요: "
          f"sha256sum {out_path.name} > {out_path.name}.sha256")


if __name__ == "__main__":
    main()
